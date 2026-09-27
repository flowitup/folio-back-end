"""French wording printed on billing document PDFs and XLSX sheets."""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.domain.billing.document import BillingDocument
from app.domain.billing.document_wording import intro_sentence
from app.domain.billing.enums import BillingDocumentKind, BillingDocumentStatus
from app.domain.billing.value_objects import BillingDocumentItem
from app.infrastructure.xlsx.billing_document_xlsx_renderer import OpenpyxlBillingDocumentXlsxRenderer


def _doc(kind: BillingDocumentKind, **overrides) -> BillingDocument:
    fields = dict(
        id=uuid4(),
        user_id=uuid4(),
        kind=kind,
        document_number="FAC-2026-001" if kind == BillingDocumentKind.FACTURE else "DEV-2026-001",
        status=BillingDocumentStatus.DRAFT,
        issue_date=date(2026, 9, 27),
        created_at=datetime(2026, 9, 27, tzinfo=timezone.utc),
        updated_at=datetime(2026, 9, 27, tzinfo=timezone.utc),
        recipient_name="Client",
        issuer_legal_name="QA-C Construction SARL",
        issuer_address="9 rue du Test, 75011 Paris",
        items=(
            BillingDocumentItem(
                description="Service", quantity=Decimal("1"), unit_price=Decimal("100"), vat_rate=Decimal("20")
            ),
        ),
    )
    fields.update(overrides)
    return BillingDocument(**fields)


def _sheet_texts(doc: BillingDocument) -> list[str]:
    ws = load_workbook(BytesIO(OpenpyxlBillingDocumentXlsxRenderer().render(doc))).active
    return [str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


@pytest.mark.parametrize(
    "kind,expected",
    [
        (BillingDocumentKind.FACTURE, "Veuillez trouver ci-après la facture relative à la mission citée en objet."),
        (BillingDocumentKind.DEVIS, "Veuillez trouver ci-après le devis relatif à la mission citée en objet."),
    ],
)
def test_intro_sentence_agrees_with_the_document_kind(kind, expected):
    assert intro_sentence(kind) == expected
    assert expected in _sheet_texts(_doc(kind))
