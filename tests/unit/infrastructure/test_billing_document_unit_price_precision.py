"""Billing PDFs and XLSX sheets print a unit price with its own decimals (15,015 €, not 15,02 €),
so a reader can multiply a line's quantity by its unit price and get the amount printed next to it."""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader

from app.domain.billing.document import BillingDocument
from app.domain.billing.enums import BillingDocumentKind, BillingDocumentStatus
from app.domain.billing.value_objects import BillingDocumentItem
from app.infrastructure.pdf.billing_document_pdf_renderer import ReportLabBillingDocumentPdfRenderer
from app.infrastructure.xlsx.billing_document_xlsx_renderer import OpenpyxlBillingDocumentXlsxRenderer


def _doc() -> BillingDocument:
    now = datetime(2026, 10, 9, tzinfo=timezone.utc)
    return BillingDocument(
        id=uuid4(),
        user_id=uuid4(),
        kind=BillingDocumentKind.DEVIS,
        document_number="DEV-2026-001",
        status=BillingDocumentStatus.DRAFT,
        issue_date=date(2026, 10, 9),
        created_at=now,
        updated_at=now,
        recipient_name="Client SARL",
        issuer_legal_name="QA Construction SARL",
        issuer_address="9 rue Test\n75011 Paris",
        validity_until=date(2026, 11, 8),
        items=(
            BillingDocumentItem("Carrelage", Decimal("3.333"), Decimal("15.015"), Decimal("20")),
            BillingDocumentItem("Plomberie", Decimal("2"), Decimal("1234.565"), Decimal("20")),
            BillingDocumentItem("Peinture", Decimal("1"), Decimal("100"), Decimal("20")),
        ),
    )


def test_pdf_prints_unit_prices_with_their_own_decimals():
    reader = PdfReader(BytesIO(ReportLabBillingDocumentPdfRenderer().render(_doc())))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    text = text.replace("\u00a0", " ").replace("\u202f", " ")
    assert "15,015 €" in text
    assert "1 234,565 €" in text
    assert "15,02 €" not in text and "1 234,57 €" not in text
    # A whole price still prints with two decimals; line amounts stay at the cent.
    assert "100,00 €" in text
    assert "50,04 €" in text and "2 469,13 €" in text


def test_xlsx_unit_price_format_shows_more_than_two_decimals():
    ws = load_workbook(BytesIO(OpenpyxlBillingDocumentXlsxRenderer().render(_doc()))).active
    prices = [c for row in ws.iter_rows(min_col=9, max_col=9) for c in row if c.value in (15.015, 1234.565, 100)]
    assert len(prices) == 3
    for cell in prices:
        assert cell.number_format == "#,##0.00####"


def test_expense_export_pdf_prints_a_released_line_unit_price_with_its_own_decimals():
    """A paid facture's lines are copied into its released-funds expense with their unit prices."""
    from app.domain.invoice.export.pdf_builder import build_pdf
    from app.domain.value_objects.invoice_item import InvoiceItem
    from tests.unit.invoice_export.test_pdf_builder import _make_bundle, _make_context, _make_invoice

    invoice = _make_invoice()
    invoice.items[:] = [InvoiceItem("Plomberie", Decimal("2"), Decimal("1234.565"), Decimal("20"))]
    reader = PdfReader(BytesIO(build_pdf(_make_context(), _make_bundle([invoice]))))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    text = text.replace("\u00a0", " ").replace("\u202f", " ")
    assert "1 234,565 €" in text
    assert "1 234,57 €" not in text
