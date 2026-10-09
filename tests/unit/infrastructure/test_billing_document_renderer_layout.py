"""What billing document PDFs and XLSX sheets print, and how the PDF lays it out."""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus.doctemplate import LayoutError

from app.domain.billing.document import BillingDocument
from app.domain.billing.enums import BillingDocumentKind, BillingDocumentStatus
from app.domain.billing.exceptions import BillingDocumentRenderError
from app.domain.billing.value_objects import BillingDocumentItem
from app.domain.entities.project import Project
from app.domain.labor.export.format import format_decimal_fr
from app.infrastructure.pdf import billing_document_pdf_renderer as pdf_module
from app.infrastructure.pdf.billing_document_pdf_renderer import ReportLabBillingDocumentPdfRenderer
from app.infrastructure.xlsx.billing_document_xlsx_renderer import OpenpyxlBillingDocumentXlsxRenderer

NOTES = "Accès chantier par la cour\nPrévoir bâche de protection\nClé chez le gardien"
SIGNATURE = "Bon pour accord, date et signature"


def _item(description="Peinture", quantity="1", unit_price="100", vat_rate="20", category=None):
    return BillingDocumentItem(
        description=description,
        quantity=Decimal(quantity),
        unit_price=Decimal(unit_price),
        vat_rate=Decimal(vat_rate),
        category=category,
    )


def _doc(kind=BillingDocumentKind.DEVIS, **overrides) -> BillingDocument:
    fields = dict(
        id=uuid4(),
        user_id=uuid4(),
        kind=kind,
        document_number="DEV-2026-009" if kind == BillingDocumentKind.DEVIS else "FAC-2026-002",
        status=BillingDocumentStatus.DRAFT,
        issue_date=date(2026, 10, 9),
        created_at=datetime(2026, 10, 9, tzinfo=timezone.utc),
        updated_at=datetime(2026, 10, 9, tzinfo=timezone.utc),
        recipient_name="Client SARL",
        recipient_address="1 a\n2 b\n3 c\n4 d\n5 e",
        recipient_email="client@example.fr",
        recipient_siret="12345678900011",
        issuer_legal_name="QA Construction SARL",
        issuer_address="9 rue Test\n75011 Paris\nFrance",
        notes=NOTES,
        terms="Acompte 30 % à la commande",
        signature_block_text=SIGNATURE,
        validity_until=date(2026, 11, 8) if kind == BillingDocumentKind.DEVIS else None,
        items=(
            _item("Peinture", "3.333", "1234.565", "10"),
            _item("Coffrage", "125.75", "1250.50", "20"),
            _item("Main d'oeuvre", "12", "1250.5", "5.5"),
        ),
    )
    fields.update(overrides)
    return BillingDocument(**fields)


def _project() -> Project:
    return Project(
        id=uuid4(),
        name="Chantier Lilas",
        owner_id=uuid4(),
        created_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
        address="5 rue des Lilas",
    )


def _pdf(doc: BillingDocument, project: Project | None = None) -> PdfReader:
    return PdfReader(BytesIO(ReportLabBillingDocumentPdfRenderer().render(doc, project=project)))


def _pdf_text(reader: PdfReader) -> str:
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return text.replace(" ", " ").replace(" ", " ")


def _sheet(doc: BillingDocument, project: Project | None = None):
    return load_workbook(BytesIO(OpenpyxlBillingDocumentXlsxRenderer().render(doc, project=project))).active


def _cells(ws) -> dict[str, object]:
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}


def _values(ws) -> list[str]:
    return [str(v).replace(" ", " ") for v in _cells(ws).values()]


# ---------------------------------------------------------------------------
# PDF — long lines (billing-04 / backend-code-06)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "description",
    ["Lorem ipsum dolor sit amet. " * 178 + "fin", "x" * 5000],
    ids=["prose", "no-spaces"],
)
def test_pdf_renders_a_line_description_longer_than_a_page(description):
    reader = _pdf(_doc(items=(_item(description),)), _project())
    assert len(reader.pages) >= 2
    text = _pdf_text(reader)
    assert "fin" in text or "xxxx" in text
    assert "Total TTC" in text


def test_pdf_renders_a_signature_and_an_address_taller_than_a_page():
    reader = _pdf(_doc(recipient_address="a\n" * 249, signature_block_text="s\n" * 249))
    assert len(reader.pages) > 2


def test_pdf_renders_several_line_descriptions_longer_than_a_page():
    """Each split used to re-centre the MIDDLE-aligned cells and grow the next tall row past a page."""
    description = ("Lorem ipsum dolor sit amet. " * 200)[:5000]
    reader = _pdf(_doc(items=(_item(description), _item(description), _item("y" * 5000))), _project())
    assert len(reader.pages) >= 4
    assert "Total TTC" in _pdf_text(reader)


def test_pdf_splits_a_row_barely_taller_than_a_page():
    """A row whose remainder would be under the 3 mm minimum still splits on an empty page."""
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    styles = pdf_module._styles()
    lines = next(
        n
        for n in range(40, 80)
        if pdf_module._FRAME_HEIGHT < n * styles["body"].leading + 6 < pdf_module._FRAME_HEIGHT + 3 * mm
    )
    table = pdf_module._PageTable([[[Paragraph("s", styles["body"]) for _ in range(lines)]]], colWidths=[200])
    table = pdf_module._split_rows_taller_than_a_page(table, 200)
    assert table.splitInRow > 0
    buf = BytesIO()
    SimpleDocTemplate(
        buf,
        topMargin=pdf_module._MARGIN,
        bottomMargin=pdf_module._BOTTOM_MARGIN,
        leftMargin=pdf_module._MARGIN,
        rightMargin=pdf_module._MARGIN,
    ).build([table])
    assert len(PdfReader(BytesIO(buf.getvalue())).pages) == 2


def test_pdf_short_lines_never_split_across_pages():
    """Only rows taller than a page may split; ordinary tables keep moving to the next page whole."""
    styles = pdf_module._styles()
    table = pdf_module._build_items_table(_doc(), styles, 500)[0]
    assert table.splitInRow == 0
    tall = pdf_module._build_items_table(_doc(items=(_item("x" * 5000),)), styles, 500)[0]
    assert tall.splitInRow > 0


def test_pdf_layout_error_becomes_a_render_error():
    with patch.object(pdf_module.SimpleDocTemplate, "build", side_effect=LayoutError("too large")):
        with pytest.raises(BillingDocumentRenderError):
            ReportLabBillingDocumentPdfRenderer().render(_doc())


# ---------------------------------------------------------------------------
# PDF — items table columns and totals (billing-15 / backend-code-07)
# ---------------------------------------------------------------------------


def test_pdf_figure_columns_fit_their_widest_value_and_header():
    doc = _doc(items=(_item("Gros oeuvre", "125.75", "1250.50", "5.5"), _item("Béton", "1", "48500", "20")))
    styles = pdf_module._styles()
    table = pdf_module._build_items_table(doc, styles, 510)[0]
    widths = table._colWidths
    padding = 8  # LEFTPADDING + RIGHTPADDING
    for col, text, font, size in [
        (2, "125,75", "DejaVu", 8),
        (3, "1 250,50 €", "DejaVu", 8),
        (3, "(HT) en €", "DejaVu-Bold", 9),
        (4, "Avancement", "DejaVu-Bold", 9),
        (5, "157 250,38 €", "DejaVu", 8),
        (5, "48 500,00 €", "DejaVu", 8),
        (6, "5,5 %", "DejaVu", 8),
        (6, "TVA", "DejaVu-Bold", 9),
    ]:
        assert widths[col] - padding >= pdfmetrics.stringWidth(text, font, size), text
    assert widths[0] >= 510 * 0.4
    assert sum(widths) == pytest.approx(510)


def test_pdf_prints_headers_and_amounts_unbroken():
    text = _pdf_text(_pdf(_doc(items=(_item("Coffrage", "125.75", "1250.50", "20"),))))
    lines = text.splitlines()
    for word in ("Avancement", "TVA", "125,75", "1 250,50 €", "157 250,38 €", "20 %"):
        assert any(word in line for line in lines), word
    assert "Avanc\n" not in text


def test_pdf_totals_are_right_aligned_like_the_ttc_row():
    totals, ttc = pdf_module._build_totals_block(_doc(), pdf_module._styles(), 510)
    for row in totals._cellvalues:
        assert all(cell.style.alignment == 2 for cell in row)
    assert all(cell.style.alignment == 2 for cell in ttc._cellvalues[0])


# ---------------------------------------------------------------------------
# PDF — footer (billing-16 / backend-code-08)
# ---------------------------------------------------------------------------


def test_pdf_footer_is_french_with_the_page_count():
    reader = _pdf(_doc(items=(_item("x" * 5000),)))
    total = len(reader.pages)
    assert total >= 2
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").replace(" ", " ")
        assert f"Généré par Folio · Page {number} / {total}" in text
    assert "Generated by" not in _pdf_text(reader)


# ---------------------------------------------------------------------------
# PDF — CJK characters (billing-17)
# ---------------------------------------------------------------------------


def test_cjk_runs_switch_to_a_cjk_font():
    assert pdf_module._pdf_text("Ñoño 日本 & co") == 'Ñoño <font name="STSong-Light">日本</font> &amp; co'
    assert pdf_module._pdf_text("テスト中文") == '<font name="STSong-Light">テスト中文</font>'
    assert pdf_module._pdf_text("한국어") == '<font name="HYGothic-Medium">한국어</font>'
    assert pdf_module._pdf_text("Nguyễn Văn Đức — 1 234,50 €") == "Nguyễn Văn Đức — 1 234,50 €"


def test_pdf_draws_cjk_text_with_the_cjk_font():
    reader = _pdf(_doc(recipient_name="Société Ñoño 日本", items=(_item("Peinture 日本語テスト"),)))
    fonts = {str(f.get_object()["/BaseFont"]) for f in reader.pages[0]["/Resources"]["/Font"].values()}
    assert any("STSong-Light" in name for name in fonts)


# ---------------------------------------------------------------------------
# PDF — notes, signature, linked project (billing-02)
# ---------------------------------------------------------------------------


def test_pdf_prints_notes_signature_and_the_linked_project():
    text = _pdf_text(_pdf(_doc(), _project()))
    assert "Notes" in text
    for line in NOTES.splitlines():
        assert text.count(line) == 1, line
    assert SIGNATURE in text
    assert "Objet : Chantier Lilas — 5 rue des Lilas" in text
    assert text.count("Chantier Lilas — 5 rue des Lilas") == 2  # Objet + project row of the table


def test_pdf_without_a_project_prints_no_objet_and_no_project_row():
    text = _pdf_text(_pdf(_doc()))
    assert "Objet :" not in text
    assert text.count("Client SARL") == 1  # no longer repeated as the table's "project"
    assert text.count("Accès chantier par la cour") == 1  # in the Notes, not as a project row


# ---------------------------------------------------------------------------
# XLSX — content parity with the PDF (billing-02 / billing-03)
# ---------------------------------------------------------------------------


def test_xlsx_prints_the_linked_project_as_objet():
    ws = _sheet(_doc(), _project())
    cells = _cells(ws)
    objet = next(coord for coord, v in cells.items() if v == "Objet/Opération")
    row = ws[objet].row
    assert cells[f"C{row + 1}"] == "Chantier Lilas"
    assert cells[f"C{row + 2}"] == "5 rue des Lilas"
    assert list(cells.values()).count("Chantier Lilas") == 2  # Objet + project row of the table


def test_xlsx_without_a_project_prints_no_objet():
    values = _values(_sheet(_doc()))
    assert "Objet/Opération" not in values
    assert values.count("Client SARL") == 1
    assert values.count("Accès chantier par la cour") == 1


def test_xlsx_prints_notes_terms_and_signature():
    values = _values(_sheet(_doc()))
    assert "Notes" in values
    for line in NOTES.splitlines():
        assert values.count(line) == 1, line
    assert "Conditions générales" in values
    assert "Acompte 30 % à la commande" in values
    assert SIGNATURE in values


def test_xlsx_keeps_user_text_starting_with_equals_as_text():
    """openpyxl would store '=> …' as a formula, which Excel reports as a corrupt file."""
    doc = _doc(
        notes="=> Accès par la cour",
        terms='=HYPERLINK("http://example.com")',
        signature_block_text="=signature",
        items=(_item("=1+1"),),
    )
    ws = _sheet(doc, _project())
    for text in ("=> Accès par la cour", '=HYPERLINK("http://example.com")', "=signature", "=1+1"):
        cell = next(c for row in ws.iter_rows() for c in row if c.value == text)
        assert cell.data_type == "s", text


def test_xlsx_prints_a_multi_line_project_address_on_one_row():
    project = _project()
    project.address = "5 rue des Lilas\n75011 Paris"
    assert "5 rue des Lilas, 75011 Paris" in _values(_sheet(_doc(), project))


def test_xlsx_devis_prints_its_validity_date():
    assert "Valide jusqu'au : 08/11/2026" in _values(_sheet(_doc()))


def test_xlsx_facture_prints_due_date_and_payment_terms():
    doc = _doc(
        BillingDocumentKind.FACTURE,
        payment_due_date=date(2026, 11, 8),
        payment_terms="Paiement à 30 jours fin de mois",
    )
    values = _values(_sheet(doc))
    assert "Échéance : 08/11/2026" in values
    assert "Conditions : Paiement à 30 jours fin de mois" in values


def test_xlsx_prints_one_vat_row_per_rate_and_sums_them_into_ttc():
    ws = _sheet(_doc())
    cells = _cells(ws)
    rows = {v.replace(" ", " "): ws[coord].row for coord, v in cells.items() if isinstance(v, str)}
    assert {"TVA 20 %", "TVA 10 %", "TVA 5,5 %"} <= set(rows)
    assert not any(ws[coord].column_letter == "H" and v == "TVA" for coord, v in cells.items())  # no lump row
    doc = _doc()
    breakdown = {f"TVA {format_decimal_fr(rate)} %": float(tva) for rate, _base, tva in doc.vat_breakdown}
    for label, tva in breakdown.items():
        assert cells[f"K{rows[label]}"] == pytest.approx(tva)
    ht, ttc = rows["Total (HT)"], rows["Total (TTC)"]
    assert cells[f"K{ttc}"] == f"=K{ht}+SUM(K{ht + 1}:K{ttc - 1})"


def test_xlsx_prints_the_whole_recipient_block():
    values = _values(_sheet(_doc()))
    for line in ("1 a", "2 b", "3 c", "4 d", "5 e", "client@example.fr", "SIRET : 12345678900011"):
        assert line in values, line


def test_xlsx_splits_the_issuer_address_over_rows():
    cells = _cells(_sheet(_doc()))
    assert cells["C2"] == "Siège : 9 rue Test"
    assert cells["C3"] == "75011 Paris"
    assert cells["C4"] == "France"
    one_line = _cells(_sheet(_doc(issuer_address="9 rue du Test, 75011 Paris")))
    assert one_line["C2"] == "Siège : 9 rue du Test, 75011 Paris"
    assert one_line["B8"] == "Réf Devis"  # a one-line address keeps the source layout
