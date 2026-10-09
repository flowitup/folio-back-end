"""The labour export renders its labels in the requested locale (en default, fr, vi)."""

from __future__ import annotations

import dataclasses
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from pypdf import PdfReader

from app.domain.labor.export.labels import _LABELS, SUPPORTED_LOCALES, month_label, range_label, shift_label, t
from app.domain.labor.export.pdf_builder import build_pdf
from app.domain.labor.export.xlsx_builder import build_xlsx
from tests.unit.domain.labor.export.test_xlsx_builder import _make_entry, _make_two_month_buckets


def _pdf_text(raw: bytes) -> str:
    return " ".join("".join(p.extract_text() or "" for p in PdfReader(BytesIO(raw)).pages).split())


def test_every_locale_has_the_same_label_keys():
    for locale in SUPPORTED_LOCALES:
        assert set(_LABELS[locale]) == set(_LABELS["en"]), locale


@pytest.mark.parametrize(
    ("locale", "expected"), [("en", "Apr 2026"), ("fr", "avr. 2026"), ("vi", "Tháng 4 năm 2026"), ("de", "Apr 2026")]
)
def test_month_label(locale, expected):
    assert month_label(date(2026, 4, 1), locale) == expected


def test_range_label_pluralises_in_english():
    assert range_label(date(2026, 4, 1), date(2026, 4, 1)).endswith("(1 month)")
    assert range_label(date(2026, 4, 1), date(2026, 5, 1)).endswith("(2 months)")


@pytest.mark.parametrize(
    ("locale", "sheets", "header"),
    [
        ("en", ["Summary", "Apr 2026", "May 2026"], "Worker"),
        ("fr", ["Synthèse", "avr. 2026", "mai 2026"], "Ouvrier"),
        ("vi", ["Tổng hợp", "Tháng 4 năm 2026", "Tháng 5 năm 2026"], "Nhân công"),
    ],
)
def test_xlsx_follows_the_locale(locale, sheets, header):
    ctx, buckets = _make_two_month_buckets()
    wb = openpyxl.load_workbook(BytesIO(build_xlsx(dataclasses.replace(ctx, locale=locale), buckets)))
    assert wb.sheetnames == sheets
    assert header in {c.value for row in wb[sheets[0]].iter_rows() for c in row}


def test_pdf_follows_the_locale():
    ctx, buckets = _make_two_month_buckets()
    text = _pdf_text(build_pdf(dataclasses.replace(ctx, locale="fr"), buckets))
    assert "Export main-d'œuvre" in text
    assert "Coût total" in text
    assert "Labor Export" not in text


@pytest.mark.parametrize(
    ("locale", "full", "overtime", "supplement"),
    [
        ("en", "Full day", "Overtime", "Extra hrs (unpaid)"),
        ("fr", "Journée complète", "Heures sup. (x1,5)", "Heures en plus (non payées)"),
        ("vi", "Cả ngày", "Tăng ca", "Giờ thêm (không tính lương)"),
    ],
)
def test_xlsx_day_log_names_shifts_in_the_locale(locale, full, overtime, supplement):
    ctx, buckets = _make_two_month_buckets()
    buckets[0].daily_entries.append(
        _make_entry(entry_date="2026-04-02", shift_type="overtime", supplement_hours=2),
    )
    buckets[0].daily_entries.append(_make_entry(entry_date="2026-04-03", shift_type=None, supplement_hours=3))
    wb = openpyxl.load_workbook(BytesIO(build_xlsx(dataclasses.replace(ctx, locale=locale), buckets)))
    values = [[c.value for c in row] for row in wb[wb.sheetnames[1]].iter_rows()]
    header = next(row for row in values if row[0] == t(locale, "date"))
    assert header[3] == supplement
    days = [row for row in values if isinstance(row[0], datetime)]
    assert [(d[0], d[2]) for d in days] == [
        (datetime(2026, 4, 1), full),
        (datetime(2026, 4, 2), overtime),
        (datetime(2026, 4, 3), None),  # supplement-only day: no shift
    ]
    # Unpaid extra hours never read as overtime ("Heures sup…") in French.
    assert not (locale == "fr" and supplement.startswith("Heures sup"))


def test_xlsx_worker_rate_is_formatted_like_the_pdf():
    ctx, buckets = _make_two_month_buckets()
    ctx = dataclasses.replace(ctx, locale="fr", worker_name="Bravo", worker_daily_rate=Decimal("100.56"))
    ws = openpyxl.load_workbook(BytesIO(build_xlsx(ctx, buckets))).active
    assert ws["A5"].value == "Ouvrier : Bravo    Tarif : 100,56\u00a0€/jour"
    assert "Tarif : 100,56 €/jour" in _pdf_text(build_pdf(ctx, buckets)).replace("\u00a0", " ")


def test_shift_label_falls_back_to_the_raw_value():
    assert shift_label("fr", "half") == "Demi-journée"
    assert shift_label("de", "half") == "Half day"
    assert shift_label("fr", None) == ""
    assert shift_label("fr", "night") == "night"
