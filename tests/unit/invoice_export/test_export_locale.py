"""The expense export renders its labels in the requested locale (en default, fr, vi)."""

from __future__ import annotations

import dataclasses
from io import BytesIO

import openpyxl
import pytest

from app.domain.invoice.export.labels import SUPPORTED_LOCALES, _LABELS, _TYPE_LABELS, t
from app.domain.invoice.export.pdf_builder import build_pdf
from app.domain.invoice.export.xlsx_builder import build_xlsx
from pypdf import PdfReader

from tests.unit.invoice_export.test_pdf_builder import _make_bundle, _make_context, _make_invoice


def _extract_text(pdf_bytes: bytes) -> str:
    return "".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf_bytes)).pages)


def test_every_locale_has_the_same_label_keys():
    for table in (_LABELS, _TYPE_LABELS):
        keys = set(table["en"])
        for locale in SUPPORTED_LOCALES:
            assert set(table[locale]) == keys, locale


def test_unknown_locale_falls_back_to_english():
    assert t("de", "invoices") == "Invoices"


@pytest.mark.parametrize(
    ("locale", "summary", "total_row"),
    [("en", "Summary", "TOTAL EXPENSES"), ("fr", "Synthèse", "TOTAL DES DÉPENSES"), ("vi", "Tổng hợp", "TỔNG CHI PHÍ")],
)
def test_xlsx_labels_follow_the_locale(locale, summary, total_row):
    ctx = dataclasses.replace(_make_context(), locale=locale)
    wb = openpyxl.load_workbook(BytesIO(build_xlsx(ctx, _make_bundle([_make_invoice()]))))
    assert wb.sheetnames[0] == summary
    values = {c.value for row in wb[summary].iter_rows() for c in row}
    assert total_row in values


def test_pdf_labels_follow_the_locale():
    ctx = dataclasses.replace(_make_context(), locale="fr")
    text = _extract_text(build_pdf(ctx, _make_bundle([_make_invoice()])))
    assert "EXPORT DES DÉPENSES" in text
    assert "Total des dépenses" in text
    assert "INVOICE EXPORT" not in text
